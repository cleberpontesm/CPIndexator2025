# app.py - VERSÃO FINAL COM CHECKBOXES PARA PREENCHIMENTO AUTOMÁTICO - CORRIGIDA
import streamlit as st
import pandas as pd
from sqlalchemy import create_engine, text
from supabase import create_client, Client
from collections import defaultdict
from io import BytesIO
import os
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

# --- Bloco de importação de bibliotecas de exportação ---
try:
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from reportlab.lib.pagesizes import A4, A3, letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table as ReportlabTable, TableStyle, Paragraph, PageBreak, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    EXPORT_LIBS_AVAILABLE = True
except ImportError:
    EXPORT_LIBS_AVAILABLE = False

# --- Definições e constantes (do seu código original) ---
FORM_DEFINITIONS = {
    "Nascimento/Batismo": ["Data do Registro", "Data do Evento", "Local do Evento", "Nome do Registrado", "Nome do Pai", "Nome da Mãe", "Padrinhos", "Avô paterno", "Avó paterna", "Avô materno", "Avó materna"],
    "Casamento": ["Data do Registro", "Data do Evento", "Local do Evento", "Nome do Noivo", "Idade do Noivo", "Pai do Noivo", "Mãe do Noivo", "Nome da Noiva", "Idade da Noiva", "Pai da Noiva", "Mãe da Noiva", "Testemunhas"],
    "Óbito": ["Data do Registro", "Data do Óbito", "Local do Óbito", "Nome do Falecido", "Idade no Óbito", "Filiação", "Cônjuge Sobrevivente", "Deixou Filhos?", "Causa Mortis", "Local do Sepultamento"],
    "Notas": ["Tipo de Ato", "Data do Registro", "Local do Registro", "Partes Envolvidas", "Resumo do Teor"]
}
COMMON_FIELDS = ["Fonte (Livro)", "Fonte (Página/Folha)", "Observações", "Caminho da Imagem"]

EXPORT_COLUMN_ORDER = {
    "Nascimento/Batismo": ["id", "tipo_registro"] + [f.lower().replace(" ", "_") for f in FORM_DEFINITIONS["Nascimento/Batismo"] + COMMON_FIELDS] + ['criado_por', 'ultima_alteracao_por', 'criado_em', 'atualizado_em'],
    "Casamento": ["id", "tipo_registro"] + [f.lower().replace(" ", "_") for f in FORM_DEFINITIONS["Casamento"] + COMMON_FIELDS] + ['criado_por', 'ultima_alteracao_por', 'criado_em', 'atualizado_em'],
    "Óbito": ["id", "tipo_registro"] + [f.lower().replace(" ", "_").replace("?", "") for f in FORM_DEFINITIONS["Óbito"] + COMMON_FIELDS] + ['criado_por', 'ultima_alteracao_por', 'criado_em', 'atualizado_em'],
    "Notas": ["id", "tipo_registro", "tipo_de_ato", "data_do_registro", "local_do_registro", "partes_envolvidas", "resumo_do_teor", "fonte_livro", "fonte_pagina_folha", "observacoes", 'criado_por', 'ultima_alteracao_por', 'criado_em', 'atualizado_em']
}

COLUMN_LABELS = {
    'id': 'ID',
    'tipo_registro': 'Tipo de Registro',
    'data_do_registro': 'Data do Registro',
    'data_do_evento': 'Data do Evento',
    'data_do_obito': 'Data do Óbito',
    'local_do_evento': 'Local do Evento',
    'local_do_obito': 'Local do Óbito',
    'nome_do_registrado': 'Nome do Registrado',
    'nome_do_pai': 'Nome do Pai',
    'nome_da_mae': 'Nome da Mãe',
    'padrinhos': 'Padrinhos',
    'avo_paterno': 'Avô Paterno',
    'avo_paterna': 'Avó Paterna',
    'avo_materno': 'Avô Materno',
    'avo_materna': 'Avó Materna',
    'nome_do_noivo': 'Nome do Noivo',
    'idade_do_noivo': 'Idade do Noivo',
    'pai_do_noivo': 'Pai do Noivo',
    'mae_do_noivo': 'Mãe do Noivo',
    'nome_da_noiva': 'Nome da Noiva',
    'idade_da_noiva': 'Idade da Noiva',
    'pai_da_noiva': 'Pai da Noiva',
    'mae_da_noiva': 'Mãe da Noiva',
    'testemunhas': 'Testemunhas',
    'nome_do_falecido': 'Nome do Falecido',
    'idade_no_obito': 'Idade no Óbito',
    'filiacao': 'Filiação',
    'conjuge_sobrevivente': 'Cônjuge Sobrevivente',
    'deixou_filhos': 'Deixou Filhos',
    'causa_mortis': 'Causa Mortis',
    'local_do_sepultamento': 'Local do Sepultamento',
    'fonte_livro': 'Fonte (Livro)',
    'fonte_pagina_folha': 'Fonte (Página/Folha)',
    'observacoes': 'Observações',
    'caminho_da_imagem': 'Caminho da Imagem',
    'criado_por': 'Criado Por',
    'ultima_alteracao_por': 'Última Alteração Por',
    'tipo_de_ato': 'Tipo de Ato',
    'local_do_registro': 'Local do Registro',
    'partes_envolvidas': 'Partes Envolvidas',
    'resumo_do_teor': 'Resumo do Teor',
    'criado_em': 'Criado Em',
    'atualizado_em': 'Atualizado Em'
}

# Lista de tipos de ato pré-definidos
TIPOS_DE_ATO = [
    "Procuração Pública",
    "Testamento",
    "Escritura de Compra e Venda",
    "Escritura de Compra e Venda de Escravos",
    "Escritura de Perfilhação/Reconhecimento de Filhos",
    "Escritura de Dote",
    "Escritura de Emancipação",
    "Escritura de Doação",
    "Escritura de Permuta",
    "Escritura de Dação em Pagamento",
    "Escritura de Usufruto",
    "Escritura Declaratória",
    "Escritura de Cessão de Direitos Hereditários",
    "Interdição",
    "Tutela",
    "Curatela",
    "Outros"
]

TABLE_COLUMNS = {
    "Nascimento/Batismo": ['id', 'nome_do_registrado', 'data_do_registro', 'data_do_evento', 'fonte_livro', 'fonte_pagina_folha'],
    "Casamento": ['id', 'nome_do_noivo', 'nome_da_noiva', 'data_do_registro', 'data_do_evento', 'fonte_livro', 'fonte_pagina_folha'],
    "Óbito": ['id', 'nome_do_falecido', 'data_do_registro', 'data_do_obito', 'fonte_livro', 'fonte_pagina_folha'],
    "Notas": ['id', 'tipo_de_ato', 'data_do_registro', 'fonte_livro', 'fonte_pagina_folha']
}

SEARCH_CATEGORIES = {
    "Nomes": [
        'nome_do_registrado', 'nome_do_pai', 'nome_da_mae', 'nome_do_noivo', 'nome_da_noiva',
        'nome_do_falecido', 'padrinhos', 'testemunhas', 'pai_do_noivo', 'mae_do_noivo',
        'pai_da_noiva', 'mae_da_noiva', 'avo_paterno', 'avo_paterna', 'avo_materno',
        'avo_materna', 'conjuge_sobrevivente', 'filiacao', 'partes_envolvidas'
    ],
    "Locais": [
        'local_do_evento', 'local_do_obito', 'local_do_registro', 'local_do_sepultamento'
    ],
    "Datas": [
        'data_do_registro', 'data_do_evento', 'data_do_obito'
    ],
    "Idades": [
        'idade_do_noivo', 'idade_da_noiva', 'idade_no_obito'
    ],
    "Informações Gerais": [
        'observacoes', 'resumo_do_teor', 'tipo_de_ato', 'causa_mortis', 'deixou_filhos', 'tipo_registro'
    ],
    "Fontes": [
        'fonte_livro', 'fonte_pagina_folha', 'caminho_da_imagem'
    ]
}

# --- CONFIGURAÇÃO INICIAL E CLIENTES ---
st.set_page_config(layout="wide", page_title="CPIndexator Web")

@st.cache_resource
def init_supabase_auth():
    try:
        url = st.secrets["SUPABASE_URL"]
        key = st.secrets["SUPABASE_KEY"]
        return create_client(url, key)
    except Exception as e:
        st.error("Erro ao inicializar o cliente Supabase. Verifique seus segredos.")
        st.stop()

@st.cache_resource
def init_db_connection():
    try:
        connection_string = st.secrets["DB_CONNECTION_STRING"]
        return create_engine(connection_string)
    except Exception as e:
        st.error("Erro ao conectar ao banco de dados. Verifique sua Connection String.")
        st.stop()

supabase = init_supabase_auth()
engine = init_db_connection()


# --- FUNÇÕES DE LÓGICA DO BANCO DE DADOS E EXPORTAÇÃO ---

def to_col_name(field_name):
    clean_name = field_name.lower().replace("ã", "a").replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u").replace("ç", "c").replace("ô", "o").replace("â", "a").replace("õ", "o")
    return clean_name.replace(" ", "_").replace("(", "").replace(")", "").replace("/", "_").replace("?", "")

@st.cache_data(ttl=300) # Cache por 5 minutos para performance
def get_distinct_values(column_name):
    with engine.connect() as conn:
        try:
            query = text(f"SELECT DISTINCT {column_name} FROM registros WHERE {column_name} IS NOT NULL AND {column_name} != '' ORDER BY {column_name}")
            result = conn.execute(query).fetchall()
            return [row[0] for row in result]
        except:
            return []

def get_table_columns():
    """Retorna as colunas existentes na tabela registros"""
    with engine.connect() as conn:
        query = text("""
            SELECT column_name
            FROM information_schema.columns
            WHERE table_name = 'registros'
        """)
        result = conn.execute(query).fetchall()
        return [row[0] for row in result]

def formatar_email_para_exibicao(email):
    """Remove a parte do domínio de uma string de e-mail para exibição."""
    if email and '@' in email:
        return email.split('@')[0]
    return email # Retorna o valor original se não for um e-mail ou for Nulo

def formatar_timestamp_para_exibicao(ts):
    """Converte um timestamp UTC para o fuso de Brasília e o formata."""
    if not ts or pd.isna(ts):  # Verifica se é None, NaT ou NaN
        return "N/D"  # Não Disponível
    
    try:
        brasilia_tz = ZoneInfo("America/Sao_Paulo")
        # Garante que o timestamp de entrada é ciente do fuso (UTC)
        if ts.tzinfo is None:
            ts = ts.replace(tzinfo=timezone.utc)
            
        local_time = ts.astimezone(brasilia_tz)
        return local_time.strftime('%d/%m/%Y %H:%M:%S')
    except (AttributeError, TypeError, ValueError):
        # Fallback caso o dado não seja um timestamp válido
        return str(ts)

def fetch_records(search_term="", selected_books=None, search_categories=None, pagina_filter=None, show_birth_parents=False, show_marriage_info=False, show_grandparents=False):
    # Nomes de exibição exatos como definidos em COLUMN_LABELS
    base_display_cols = ['ID', 'Tipo de Registro', 'Data', 'Nome Principal', 'Fonte (Livro)']
    meta_display_cols = ['Fonte (Página/Folha)', 'Criado Por', 'Criado Em', 'Última Alteração Por', 'Atualizado Em']

    # Colunas opcionais que podem ser adicionadas
    optional_display_cols = []
    if show_birth_parents:
        optional_display_cols.extend(['Nome do Pai', 'Nome da Mãe'])
    if show_marriage_info:
        optional_display_cols.extend(['Nome da Noiva', 'Pai do Noivo', 'Mãe do Noivo', 'Pai da Noiva', 'Mãe da Noiva'])
    if show_grandparents:
        optional_display_cols.extend(['Avô Paterno', 'Avó Paterna', 'Avô Materno', 'Avó Materna'])

    # Junta todas as colunas possíveis para o caso de um DataFrame vazio
    all_possible_display_cols = base_display_cols + list(dict.fromkeys(optional_display_cols)) + meta_display_cols

    if not selected_books:
        return pd.DataFrame(columns=all_possible_display_cols)

    try:
        with engine.connect() as conn:
            # A lógica de construção da query SQL permanece a mesma
            base_query = "SELECT * FROM registros WHERE fonte_livro = ANY(:books)"
            params = {'books': selected_books}

            if pagina_filter:
                base_query += " AND CAST(fonte_pagina_folha AS TEXT) ILIKE :pagina"
                params['pagina'] = f'%{pagina_filter}%'

            order_clause = " ORDER BY fonte_livro, NULLIF(regexp_replace(fonte_pagina_folha, '[^0-9].*$', ''), '')::integer NULLS LAST, fonte_pagina_folha"

            if search_term:
                search_fields = []
                if search_categories and len(search_categories) > 0:
                    for category in search_categories:
                        if category in SEARCH_CATEGORIES:
                            search_fields.extend(SEARCH_CATEGORIES[category])
                    search_fields = list(set(search_fields))
                else:
                    for fields_list in SEARCH_CATEGORIES.values():
                        search_fields.extend(fields_list)
                    search_fields.extend(['id', 'criado_por', 'ultima_alteracao_por'])
                    search_fields = list(set(search_fields))

                search_conditions = []
                for field in search_fields:
                    if field == 'id':
                        search_conditions.append(f"CAST({field} AS TEXT) ILIKE :search_term")
                    else:
                        search_conditions.append(f"COALESCE({field}, '') ILIKE :search_term")
                
                if search_conditions:
                    search_logic = f" AND ({' OR '.join(search_conditions)})"
                    query = base_query + search_logic + order_clause
                    params['search_term'] = f'%{search_term}%'
                else:
                    query = base_query + order_clause
            else:
                query = base_query + order_clause

            result = conn.execute(text(query), params)
            df = pd.DataFrame(result.fetchall())

            if not df.empty:
                df.columns = result.keys()

                # 1. Preenche colunas de dados consolidados
                df['Nome Principal'] = df.apply(
                    lambda row: row.get('nome_do_noivo') if row['tipo_registro'] == 'Casamento' else
                                row.get('nome_do_registrado') if row['tipo_registro'] == 'Nascimento/Batismo' else
                                row.get('nome_do_falecido') if row['tipo_registro'] == 'Óbito' else
                                str(row.get('partes_envolvidas', 'N/A')).split(';')[0],
                    axis=1
                )
                df['Data'] = df.apply(
                    lambda row: row.get('data_do_evento') or row.get('data_do_obito') or row.get('data_do_registro'),
                    axis=1
                )
                
                # 2. Renomeia TODAS as colunas do banco para os nomes de exibição
                df.rename(columns=COLUMN_LABELS, inplace=True)

                # 3. Formata os dados nas colunas já renomeadas
                if 'Criado Por' in df.columns:
                    df['Criado Por'] = df['Criado Por'].apply(formatar_email_para_exibicao)
                if 'Última Alteração Por' in df.columns:
                    df['Última Alteração Por'] = df['Última Alteração Por'].apply(formatar_email_para_exibicao)
                if 'Criado Em' in df.columns:
                    df['Criado Em'] = pd.to_datetime(df['Criado Em'], errors='coerce').apply(formatar_timestamp_para_exibicao)
                if 'Atualizado Em' in df.columns:
                    df['Atualizado Em'] = pd.to_datetime(df['Atualizado Em'], errors='coerce').apply(formatar_timestamp_para_exibicao)

                # 4. Define a lista final de colunas a serem exibidas
                final_display_cols = base_display_cols + optional_display_cols + meta_display_cols
                
                # 5. Filtra o DataFrame para mostrar apenas as colunas desejadas que realmente existem
                # Remove duplicatas e garante que a coluna está no DataFrame antes de tentar exibi-la
                cols_to_render = [col for col in list(dict.fromkeys(final_display_cols)) if col in df.columns]
                
                return df[cols_to_render]
            else:
                return pd.DataFrame(columns=all_possible_display_cols)

    except Exception as e:
        st.error(f"Erro ao buscar registros: {str(e)}")
        st.info("Verifique se a estrutura do banco de dados está correta.")
        return pd.DataFrame(columns=all_possible_display_cols)


def fetch_single_record(record_id):
    with engine.connect() as conn:
        query = text("SELECT * FROM registros WHERE id = :id")
        result = conn.execute(query, {'id': record_id}).fetchone()
        return result._asdict() if result else None

def generate_excel_bytes(records_by_type):
    if not EXPORT_LIBS_AVAILABLE:
        st.error("Bibliotecas de exportação não disponíveis.")
        return None
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for record_type, records in records_by_type.items():
            if records:
                df = pd.DataFrame(records)
                if record_type in EXPORT_COLUMN_ORDER:
                    columns_order = [col for col in EXPORT_COLUMN_ORDER[record_type] if col in df.columns]
                    df = df[columns_order]

                sheet_name = record_type.replace("/", "-")[:31]
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.sheets[sheet_name]
                for column in worksheet.columns:
                    max_length = 0
                    column = [cell for cell in column]
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
                        except: pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column[0].column_letter].width = adjusted_width
    return output.getvalue()

def generate_pdf_table(records_by_type):
    if not EXPORT_LIBS_AVAILABLE: st.error("Bibliotecas de exportação não disponíveis."); return None
    output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=landscape(A3)); story = []; styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=24, textColor=colors.HexColor('#1f4788'), spaceAfter=30, alignment=TA_CENTER)
    section_style = ParagraphStyle('SectionTitle', parent=styles['Heading2'], fontSize=18, textColor=colors.HexColor('#2e5090'), spaceAfter=20)
    story.append(Paragraph("Índice de Registros - CPIndexator", title_style)); story.append(Spacer(1, 0.5*inch))
    for record_type, records in records_by_type.items():
        if records:
            story.append(Paragraph(f"Registros de {record_type}", section_style)); story.append(Paragraph(f"Total: {len(records)} registros", styles['Normal'])); story.append(Spacer(1, 0.2*inch))
            df = pd.DataFrame(records)
            if record_type in TABLE_COLUMNS:
                columns_to_show = [col for col in TABLE_COLUMNS[record_type] if col in df.columns]
            else:
                columns_to_show = [col for col in df.columns if col not in ['criado_por', 'ultima_alteracao_por', 'caminho_da_imagem', 'criado_em', 'atualizado_em']]
            df_filtered = df[columns_to_show]
            headers = [COLUMN_LABELS.get(col, col.replace('_', ' ').title()) for col in columns_to_show]
            data = [headers]
            for _, row in df_filtered.iterrows():
                row_data = []
                for col in columns_to_show:
                    value_raw = row[col]
                    if col == 'fonte_pagina_folha':
                        value = str(value_raw) if pd.notna(value_raw) else '—'
                    else:
                        value = str(value_raw) if pd.notna(value_raw) else ''

                    if len(value) > 50: value = value[:47] + '...'
                    if col == 'partes_envolvidas': value = value.replace(';', ', ')
                    row_data.append(value)
                data.append(row_data)
            table = ReportlabTable(data);
            table_style = TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 11), ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'), ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('GRID', (0, 0), (-1, -1), 1, colors.black), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f0f0')]),
                ('BOTTOMPADDING', (0, 1), (-1, -1), 6), ('TOPPADDING', (0, 1), (-1, -1), 6),
            ])
            table.setStyle(table_style); story.append(table); story.append(PageBreak())
    doc.build(story); return output.getvalue()

def generate_pdf_detailed(records_by_type):
    if not EXPORT_LIBS_AVAILABLE: st.error("Bibliotecas de exportação não disponíveis."); return None
    output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18); story = []; styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=24, textColor=colors.HexColor('#1f4788'), spaceAfter=30, alignment=TA_CENTER)
    section_style = ParagraphStyle('SectionTitle', parent=styles['Heading2'], fontSize=18, textColor=colors.HexColor('#2e5090'), spaceAfter=20, spaceBefore=30)
    record_header_style = ParagraphStyle('RecordHeader', parent=styles['Heading3'], fontSize=14, textColor=colors.HexColor('#333333'), spaceAfter=12, leftIndent=20)
    field_style = ParagraphStyle('FieldStyle', parent=styles['Normal'], fontSize=11, leftIndent=40, spaceAfter=8)
    story.append(Paragraph("Relatório Detalhado de Registros - CPIndexator", title_style)); story.append(Spacer(1, 0.5*inch))
    for record_type, records in records_by_type.items():
        if records:
            story.append(Paragraph(f"Registros de {record_type}", section_style)); story.append(Paragraph(f"Total de registros: {len(records)}", styles['Normal'])); story.append(Spacer(1, 0.2*inch))
            for idx, record in enumerate(records, 1):
                nome_principal = record.get('nome_do_registrado') or record.get('nome_do_noivo') or record.get('nome_do_falecido') or str(record.get('partes_envolvidas', 'N/A')).split(';')[0] or 'Sem nome'
                header_text = f"Registro #{idx} - ID: {record.get('id', 'N/A')} - {nome_principal}"
                story.append(Paragraph(header_text, record_header_style))
                data = []
                fields_order = sorted(record.keys())
                if record_type in EXPORT_COLUMN_ORDER:
                    fields_order = [col for col in EXPORT_COLUMN_ORDER[record_type] if col in record]

                for field in fields_order:
                    if field in record and ((pd.notna(record[field]) and record[field] != '') or field == 'fonte_pagina_folha'):
                        label = COLUMN_LABELS.get(field, field.replace('_', ' ').title())
                        value_raw = record[field]
                        
                        if field == 'fonte_pagina_folha':
                            value = str(value_raw) if pd.notna(value_raw) and value_raw != '' else '—'
                        elif field in ['criado_por', 'ultima_alteracao_por']:
                            value = formatar_email_para_exibicao(str(value_raw))
                        elif field in ['criado_em', 'atualizado_em']:
                            value = formatar_timestamp_para_exibicao(pd.to_datetime(value_raw, errors='coerce'))
                        elif field == 'partes_envolvidas':
                            value = str(value_raw).replace(';', '<br/>- ')
                            value = f"- {value}"
                        else:
                            value = str(value_raw)
                        
                        if len(value) > 60:
                            value = Paragraph(value, styles['Normal'])
                        
                        data.append([Paragraph(f"<b>{label}:</b>", field_style), value])

                if not data: data.append([Paragraph("Sem dados disponíveis", field_style), ""])
                table = ReportlabTable(data, colWidths=[2.5*inch, 4*inch]);
                table.setStyle(TableStyle([
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'), ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6), ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey), ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f0f0f0')),
                ]));
                story.append(table); story.append(Spacer(1, 0.3*inch))
                if idx < len(records): story.append(Paragraph("<hr/>", styles['Normal'])); story.append(Spacer(1, 0.1*inch))
            story.append(PageBreak())
    doc.build(story); return output.getvalue()


# --- INTERFACE DO APLICATIVO ---

def login_form():
    st.title("CPIndexator - Versão WEB")
    st.subheader("Por favor, faça o login para continuar")
    with st.form("login_form"):
        email = st.text_input("Email")
        password = st.text_input("Senha", type="password")
        submitted = st.form_submit_button("Entrar")
        if submitted:
            try:
                response = supabase.auth.sign_in_with_password({"email": email, "password": password})
                st.session_state.user = response.user
                st.rerun()
            except Exception as e:
                st.error("Email ou senha inválidos.")

def main_app():
    st.sidebar.title("Bem-vindo(a)!")
    st.sidebar.info(f"Logado como: {formatar_email_para_exibicao(st.session_state.user.email)}")
    if st.sidebar.button("Sair (Logout)"):
        for key in list(st.session_state.keys()):
            del st.session_state[key]
        st.cache_data.clear()
        st.cache_resource.clear()
        st.rerun()

    st.title("CPIndexator - Painel Principal")

    user_email = ""
    if hasattr(st.session_state, 'user') and st.session_state.user is not None:
        user_email = st.session_state.user.email
    admin_list = st.secrets.get("ADMIN_USERS", [])
    is_admin = user_email in admin_list

    # --- INÍCIO DA GRANDE MUDANÇA: SISTEMA DE ABAS MANUAL ---

    # Define as abas disponíveis
    tabs = ["➕ Adicionar Registro", "🔍 Consultar e Gerenciar", "📤 Exportar Dados"]
    if is_admin:
        tabs.append("⚙️ Administração")

    # Inicializa a aba ativa no session_state se ela não existir
    if 'active_tab' not in st.session_state:
        st.session_state.active_tab = tabs[0] # Começa na primeira aba

    # Função de callback para atualizar a aba ativa
    def set_active_tab():
        st.session_state.active_tab = st.session_state.radio_tabs

    # Usa st.radio com aparência de abas para controlar o estado
    st.radio(
        label="Navegação",
        options=tabs,
        key="radio_tabs", # Chave para o session_state
        on_change=set_active_tab, # Callback para atualizar nosso estado manual
        horizontal=True,
        label_visibility="collapsed" # Esconde o rótulo "Navegação"
    )

    # Renderiza o conteúdo da aba com base no estado que controlamos
    # Em vez de 'with tab_add:', usamos 'if st.session_state.active_tab == "nome_da_aba":'
    
    if st.session_state.active_tab == "➕ Adicionar Registro":
        st.header("Adicionar Novo Registro")
        
        # Inicializa as variáveis de sessão para os checkboxes
        if 'fixar_livro' not in st.session_state:
            st.session_state.fixar_livro = False
        if 'fixar_local' not in st.session_state:
            st.session_state.fixar_local = False
        if 'livro_fixo' not in st.session_state:
            st.session_state.livro_fixo = ""
        if 'local_fixo' not in st.session_state:
            st.session_state.local_fixo = ""

        record_type = st.selectbox("Tipo de Registro:", list(FORM_DEFINITIONS.keys()), index=None, placeholder="Selecione...")

        if 'current_record_type' not in st.session_state or st.session_state.current_record_type != record_type:
            st.session_state.current_record_type = record_type
            if 'num_partes' in st.session_state:
                del st.session_state.num_partes

        if record_type:
            if record_type == "Notas":
                # Controle dinâmico de partes envolvidas FORA do formulário
                if 'num_partes' not in st.session_state:
                    st.session_state.num_partes = 2
                
                # Interface para controlar número de partes envolvidas
                st.markdown("### Configurar Partes Envolvidas")
                col_info, col_btn1, col_btn2 = st.columns([2, 1, 1])
                with col_info:
                    st.info(f"Atualmente configurado para {st.session_state.num_partes} partes envolvidas")
                with col_btn1:
                    if st.button("➕ Adicionar Parte", key="add_parte_btn"):
                        st.session_state.num_partes += 1
                        st.rerun()
                with col_btn2:
                    if st.button("➖ Remover Última", key="remove_parte_btn") and st.session_state.num_partes > 1:
                        st.session_state.num_partes -= 1
                        st.rerun()
                
                st.markdown("---")
                
                with st.form("new_record_form", clear_on_submit=True):
                    entries = {}
                    fields = FORM_DEFINITIONS.get(record_type, []) + COMMON_FIELDS
                    
                    # Renderizar campos até "Partes Envolvidas"
                    partes_index = fields.index("Partes Envolvidas")
                    for field in fields[:partes_index]:
                        if field == "Fonte (Livro)":
                            col1, col2 = st.columns([4, 1])
                            with col1:
                                entries[to_col_name(field)] = st.text_input(
                                    f"{field}:", 
                                    value=st.session_state.livro_fixo if st.session_state.fixar_livro else "",
                                    key=f"add_{to_col_name(field)}"
                                )
                            with col2:
                                st.session_state.fixar_livro = st.checkbox(
                                    "Fixar", 
                                    value=st.session_state.fixar_livro,
                                    key=f"fixar_{to_col_name(field)}",
                                    help="Marque para usar este valor automaticamente nos próximos registros"
                                )
                                if st.session_state.fixar_livro and entries[to_col_name(field)]:
                                    st.session_state.livro_fixo = entries[to_col_name(field)]
                        elif field in ["Local do Evento", "Local do Registro"]:
                            col1, col2 = st.columns([4, 1])
                            with col1:
                                entries[to_col_name(field)] = st.text_input(
                                    f"{field}:", 
                                    value=st.session_state.local_fixo if st.session_state.fixar_local else "",
                                    key=f"add_{to_col_name(field)}"
                                )
                            with col2:
                                st.session_state.fixar_local = st.checkbox(
                                    "Fixar", 
                                    value=st.session_state.fixar_local,
                                    key=f"fixar_{to_col_name(field)}",
                                    help="Marque para usar este valor automaticamente nos próximos registros"
                                )
                                if st.session_state.fixar_local and entries[to_col_name(field)]:
                                    st.session_state.local_fixo = entries[to_col_name(field)]
                        elif field == "Tipo de Ato":
                            entries[to_col_name(field)] = st.selectbox(
                                f"{field}:",
                                options=TIPOS_DE_ATO,
                                index=None,
                                placeholder="Selecione um tipo...",
                                key=f"add_{to_col_name(field)}"
                            )
                        elif field in ["Data do Registro", "Data do Evento", "Data do Óbito"]:
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:", 
                                placeholder="DD/MM/AAAA",
                                key=f"add_{to_col_name(field)}"
                            )
                        elif field == "Fonte (Página/Folha)":
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:",
                                help="Exemplos: '12', '15v' (para verso), '34-36' (para intervalos)",
                                key=f"add_{to_col_name(field)}"
                            )
                        else:
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:", 
                                key=f"add_{to_col_name(field)}"
                            )

                    # Campos dinâmicos para "Partes Envolvidas"
                    st.markdown("#### Partes Envolvidas 👥")
                    partes_envolvidas_inputs = []
                    for i in range(st.session_state.get('num_partes', 2)):
                        partes_envolvidas_inputs.append(
                            st.text_input(f"Parte Envolvida {i+1}", key=f"add_parte_{i}")
                        )

                    # Campos restantes após "Partes Envolvidas"
                    st.markdown("---")
                    for field in fields[partes_index+1:]:
                        if field == "Fonte (Livro)":
                            col1, col2 = st.columns([4, 1])
                            with col1:
                                entries[to_col_name(field)] = st.text_input(
                                    f"{field}:", 
                                    value=st.session_state.livro_fixo if st.session_state.fixar_livro else "",
                                    key=f"add_{to_col_name(field)}"
                                )
                            with col2:
                                st.session_state.fixar_livro = st.checkbox(
                                    "Fixar", 
                                    value=st.session_state.fixar_livro,
                                    key=f"fixar_{to_col_name(field)}",
                                    help="Marque para usar este valor automaticamente nos próximos registros"
                                )
                                if st.session_state.fixar_livro and entries[to_col_name(field)]:
                                    st.session_state.livro_fixo = entries[to_col_name(field)]
                        elif field == "Fonte (Página/Folha)":
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:",
                                help="Exemplos: '12', '15v' (para verso), '34-36' (para intervalos)",
                                key=f"add_{to_col_name(field)}"
                            )
                        else:
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:", 
                                key=f"add_{to_col_name(field)}"
                            )

                    submitted = st.form_submit_button(f"Adicionar Registro de {record_type}")
                    
                    if submitted:
                        # Processar partes envolvidas
                        partes_values = [p.strip() for p in partes_envolvidas_inputs if p.strip()]
                        entries['partes_envolvidas'] = "; ".join(partes_values)

                        try:
                            with engine.connect() as conn:
                                now_utc = datetime.now(timezone.utc)
                                
                                cols = ["tipo_registro"] + list(entries.keys()) + ["criado_por", "ultima_alteracao_por", "criado_em", "atualizado_em"]
                                vals = [record_type] + list(entries.values()) + [user_email, user_email, now_utc, now_utc]

                                final_cols = []
                                final_vals = []
                                seen_cols = set()
                                for c, v in zip(cols, vals):
                                    if c not in seen_cols:
                                        final_cols.append(c)
                                        final_vals.append(v)
                                        seen_cols.add(c)

                                placeholders = ', '.join([f':{c}' for c in final_cols])
                                query = f"INSERT INTO registros ({', '.join(final_cols)}) VALUES ({placeholders})"
                                params = dict(zip(final_cols, final_vals))

                                conn.execute(text(query), params)
                                conn.commit()
                                st.success("Registro adicionado com sucesso!")
                                st.cache_data.clear() # Limpa o cache para atualizar os filtros
                                st.cache_resource.clear()
                                if 'num_partes' in st.session_state:
                                    del st.session_state.num_partes
                                st.rerun()
                        except Exception as e:
                            st.error(f"Ocorreu um erro ao salvar: {e}")
            else:
                # Para outros tipos de registro (sem partes envolvidas dinâmicas)
                with st.form("new_record_form", clear_on_submit=True):
                    entries = {}
                    fields = FORM_DEFINITIONS.get(record_type, []) + COMMON_FIELDS
                    
                    for field in fields:
                        if field == "Fonte (Livro)":
                            col1, col2 = st.columns([4, 1])
                            with col1:
                                entries[to_col_name(field)] = st.text_input(
                                    f"{field}:", 
                                    value=st.session_state.livro_fixo if st.session_state.fixar_livro else "",
                                    key=f"add_{to_col_name(field)}"
                                )
                            with col2:
                                st.session_state.fixar_livro = st.checkbox(
                                    "Fixar", 
                                    value=st.session_state.fixar_livro,
                                    key=f"fixar_{to_col_name(field)}",
                                    help="Marque para usar este valor automaticamente nos próximos registros"
                                )
                                if st.session_state.fixar_livro and entries[to_col_name(field)]:
                                    st.session_state.livro_fixo = entries[to_col_name(field)]
                        elif field in ["Local do Evento", "Local do Registro"]:
                            col1, col2 = st.columns([4, 1])
                            with col1:
                                entries[to_col_name(field)] = st.text_input(
                                    f"{field}:", 
                                    value=st.session_state.local_fixo if st.session_state.fixar_local else "",
                                    key=f"add_{to_col_name(field)}"
                                )
                            with col2:
                                st.session_state.fixar_local = st.checkbox(
                                    "Fixar", 
                                    value=st.session_state.fixar_local,
                                    key=f"fixar_{to_col_name(field)}",
                                    help="Marque para usar este valor automaticamente nos próximos registros"
                                )
                                if st.session_state.fixar_local and entries[to_col_name(field)]:
                                    st.session_state.local_fixo = entries[to_col_name(field)]
                        elif field == "Tipo de Ato":
                            entries[to_col_name(field)] = st.selectbox(
                                f"{field}:",
                                options=TIPOS_DE_ATO,
                                index=None,
                                placeholder="Selecione um tipo...",
                                key=f"add_{to_col_name(field)}"
                            )
                        elif field in ["Data do Registro", "Data do Evento", "Data do Óbito"]:
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:", 
                                placeholder="DD/MM/AAAA",
                                key=f"add_{to_col_name(field)}"
                            )
                        elif field == "Fonte (Página/Folha)":
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:",
                                help="Exemplos: '12', '15v' (para verso), '34-36' (para intervalos)",
                                key=f"add_{to_col_name(field)}"
                            )
                        else:
                            entries[to_col_name(field)] = st.text_input(
                                f"{field}:", 
                                key=f"add_{to_col_name(field)}"
                            )

                    submitted = st.form_submit_button(f"Adicionar Registro de {record_type}")
                    
                    if submitted:
                        try:
                            with engine.connect() as conn:
                                now_utc = datetime.now(timezone.utc)
                                
                                cols = ["tipo_registro"] + list(entries.keys()) + ["criado_por", "ultima_alteracao_por", "criado_em", "atualizado_em"]
                                vals = [record_type] + list(entries.values()) + [user_email, user_email, now_utc, now_utc]

                                final_cols = []
                                final_vals = []
                                seen_cols = set()
                                for c, v in zip(cols, vals):
                                    if c not in seen_cols:
                                        final_cols.append(c)
                                        final_vals.append(v)
                                        seen_cols.add(c)

                                placeholders = ', '.join([f':{c}' for c in final_cols])
                                query = f"INSERT INTO registros ({', '.join(final_cols)}) VALUES ({placeholders})"
                                params = dict(zip(final_cols, final_vals))

                                conn.execute(text(query), params)
                                conn.commit()
                                st.success("Registro adicionado com sucesso!")
                                st.cache_data.clear() # Limpa o cache para atualizar os filtros
                                st.cache_resource.clear()
                                st.rerun()
                        except Exception as e:
                            st.error(f"Ocorreu um erro ao salvar: {e}")
                            
    elif st.session_state.active_tab == "🔍 Consultar e Gerenciar":
        st.header("Consultar Registros")

        # --- INÍCIO DA MODIFICAÇÃO ---
        # As opções de visualização foram movidas da sidebar para o corpo principal da aba.
        # Usamos colunas para um layout mais organizado.
        st.markdown("---")
        st.caption("Opções de Visualização da Tabela")
        col_vis1, col_vis2, col_vis3 = st.columns(3)
        with col_vis1:
            show_birth_parents = st.checkbox(
                "Exibir Pais (Nasc./Bat.)", 
                value=False, 
                key="show_birth_parents",
                help="Marque para mostrar as colunas 'Nome do Pai' e 'Nome da Mãe'."
            )
        with col_vis2:
            show_marriage_info = st.checkbox(
                "Exibir Noiva e Pais (Cas.)", 
                value=False, 
                key="show_marriage_info",
                help="Marque para mostrar as colunas com os dados da noiva e dos pais dos noivos."
            )
        with col_vis3:
            show_grandparents = st.checkbox(
                "Exibir Avós (Nasc./Bat.)", 
                value=False, 
                key="show_grandparents",
                help="Marque para mostrar as colunas dos avós."
            )
        st.markdown("---")
        # --- FIM DA MODIFICAÇÃO ---

        # A lógica da sidebar para os filtros de busca permanece a mesma
        st.sidebar.header("Filtros de Consulta")
        all_books_manage = get_distinct_values("fonte_livro")

        if not all_books_manage:
            st.warning("Nenhum livro encontrado no banco de dados. Adicione registros primeiro.")
            selected_books_manage = []
        else:
            # ... (O resto do código da sidebar para selecionar livros, etc., continua aqui sem alterações)
            col_select_all, _ = st.sidebar.columns([1, 3])
            with col_select_all:
                if st.button("Selecionar Todos", key="select_all_books", use_container_width=True):
                    st.session_state.manage_books_select = all_books_manage
                    st.rerun()
                
                st.markdown("""
                <style>
                    div[data-testid="stHorizontalBlock"] > div:first-child button {
                        font-size: 10px !important;
                        padding: 0.25em 0.5em !important;
                    }
                </style>
                """, unsafe_allow_html=True)
                        
            selected_books_manage = st.sidebar.multiselect(
                "Filtrar por Livro(s):", 
                all_books_manage, 
                default=all_books_manage, 
                key="manage_books_select"
            )

        pagina_filter = st.sidebar.text_input("Filtrar por página/folha:", help="Busca por parte do número da folha/página. Ex: '15' encontrará '15', '15v', etc.")
        st.sidebar.subheader("🔍 Busca Avançada")
        search_term = st.sidebar.text_input("Termo de Busca:", help="Digite qualquer palavra ou frase que deseja encontrar")
        search_categories = st.sidebar.multiselect(
            "Buscar nas Categorias:",
            options=list(SEARCH_CATEGORIES.keys()),
            default=[],
            help="Selecione em quais tipos de informação buscar. Se nenhuma categoria for selecionada, a busca será feita em todos os campos.",
            key="search_categories_select"
        )
        
        # ... (O restante do código da sidebar e da aba continua exatamente como estava antes)
        # É crucial que você REMOVA as linhas antigas que estavam na sidebar:
        # st.sidebar.subheader("Opções de Visualização da Tabela") -> REMOVER
        # show_parents = st.sidebar.checkbox(...) -> REMOVER
        # show_grandparents = st.sidebar.checkbox(...) -> REMOVER

        if search_categories:
            st.sidebar.info(f"🎯 Buscando apenas em: {', '.join(search_categories)}")
            with st.sidebar.expander("Ver campos incluídos"):
                for category in search_categories:
                    if category in SEARCH_CATEGORIES:
                        fields = [COLUMN_LABELS.get(field, field.replace('_', ' ').title()) for field in SEARCH_CATEGORIES[category]]
                        st.write(f"**{category}:** {', '.join(fields)}")
        else:
            st.sidebar.info("🌐 Buscando em todos os campos disponíveis")

        with st.sidebar.expander("ℹ️ Como usar a Busca Avançada"):
            st.markdown("""
            **Busca por Termo:**
            - Digite qualquer palavra ou frase
            - A busca não diferencia maiúsculas/minúsculas
            - Use termos parciais (ex: "Fort" encontra "Fortaleza")
            
            **Filtro por Categorias:**
            - **Nomes**: Busca em todos os campos de nomes de pessoas
            - **Locais**: Busca apenas em campos de localização
            - **Datas**: Busca em todos os campos de data
            - **Idades**: Busca em campos de idade
            - **Informações Gerais**: Observações, resumos, tipos de ato
            - **Fontes**: Livros, páginas, caminhos de imagem
            
            **Dicas:**
            - Deixe as categorias vazias para buscar em tudo
            - Combine termo + categoria para busca específica
            - Use o botão "Limpar Filtros" para resetar
            """)

        if not selected_books_manage:
            st.warning("Por favor, selecione ao menos um livro no filtro.")
        else:
            import time
            start_time = time.time()
            df_records = fetch_records(search_term, selected_books_manage, search_categories, pagina_filter, show_birth_parents=show_birth_parents, show_marriage_info=show_marriage_info, show_grandparents=show_grandparents)
            search_time = time.time() - start_time
            
            if not df_records.empty:
                total_results = len(df_records)
                if search_term:
                    if search_categories:
                        st.success(f"📊 Encontrados **{total_results}** registros contendo **'{search_term}'** nas categorias: {', '.join(search_categories)} ⏱️ ({search_time:.2f}s)")
                    else:
                        st.success(f"📊 Encontrados **{total_results}** registros contendo **'{search_term}'** em qualquer campo ⏱️ ({search_time:.2f}s)")
                else:
                    st.info(f"📊 Exibindo **{total_results}** registros dos livros selecionados ⏱️ ({search_time:.2f}s)")
                
                if 'Tipo de Registro' in df_records.columns:
                    tipo_counts = df_records['Tipo de Registro'].value_counts()
                    stats_text = " | ".join([f"{tipo}: {count}" for tipo, count in tipo_counts.items()])
                    st.caption(f"📈 Distribuição por tipo: {stats_text}")
                    
            else:
                if search_term:
                    st.warning(f"❌ Nenhum registro encontrado para o termo **'{search_term}'** ⏱️ ({search_time:.2f}s)")
                    if search_categories:
                        st.info("💡 Tente expandir as categorias de busca ou remover filtros para ver mais resultados")
                else:
                    st.info("📋 Nenhum registro encontrado nos livros selecionados")
            
            st.dataframe(df_records, use_container_width=True, hide_index=True)

            if search_term or search_categories:
                st.sidebar.markdown("---")
                if st.sidebar.button("🗑️ Limpar Filtros de Busca"):
                    st.session_state.search_categories_select = []
                    st.rerun()

            st.markdown("---")
            st.header("Gerenciar Registro Selecionado")
            # O restante do código de gerenciamento (editar, excluir, etc.) continua o mesmo...
            record_id_to_manage = st.number_input(
                "Digite o ID do registro e tecle ENTER para ver detalhes, editar ou excluir:", 
                min_value=1, 
                step=1, 
                value=None, 
                key="record_id_input"
            )

            if record_id_to_manage:
                if 'record_id' not in st.session_state or st.session_state.record_id != record_id_to_manage:
                    st.session_state.record_id = record_id_to_manage
                    if 'manage_action' in st.session_state: 
                        del st.session_state.manage_action
                    if 'edit_num_partes' in st.session_state: 
                        del st.session_state.edit_num_partes

                record = fetch_single_record(record_id_to_manage)
                if record:
                    col1, col2, col3 = st.columns(3)
                    with col1:
                        if st.button("📋 Ver Detalhes", use_container_width=True): 
                            st.session_state.manage_action = "view"
                            st.rerun()
                    with col2:
                        if st.button("✏️ Editar", use_container_width=True): 
                            st.session_state.manage_action = "edit"
                            st.rerun()
                    with col3:
                        if st.button("🗑️ Excluir", use_container_width=True): 
                            st.session_state.manage_action = "delete"
                            st.rerun()
                else: 
                    st.error(f"Registro ID {record_id_to_manage} não encontrado.")

            if 'manage_action' in st.session_state and 'record_id' in st.session_state and st.session_state.record_id:
                record_id = st.session_state.record_id
                record = fetch_single_record(record_id)
                if not record: 
                    st.error(f"Registro ID {record_id} não encontrado. Pode ter sido excluído.")
                    return

                action = st.session_state.manage_action
                st.subheader(f"Ação: {action.title()} | Registro ID: {record_id}")
                st.markdown("---")
                
                if action == "view":
                    for key, value in record.items():
                        if value:
                            label = COLUMN_LABELS.get(key, key.replace('_', ' ').title())
                            display_value = value

                            if key in ['criado_por', 'ultima_alteracao_por']:
                                display_value = formatar_email_para_exibicao(value)
                            elif key in ['criado_em', 'atualizado_em']:
                                display_value = formatar_timestamp_para_exibicao(pd.to_datetime(value, errors='coerce'))
                            elif key == 'partes_envolvidas':
                                display_value = str(value).replace(';', ' | ')
                            
                            st.write(f"**{label}:** {display_value}")

                elif action == "edit":
                    record_type = record.get('tipo_registro')
                    if not record_type: 
                        st.error("Tipo de registro não definido. Não é possível editar.")
                        return

                    if record_type == "Notas":
                        partes_str = record.get('partes_envolvidas', '')
                        partes_list = partes_str.split('; ') if partes_str else []

                        if 'edit_num_partes' not in st.session_state:
                            st.session_state.edit_num_partes = max(1, len(partes_list))

                        # Controle dinâmico de partes envolvidas para edição
                        st.markdown("### Configurar Partes Envolvidas (Edição)")
                        col_info, col_btn1, col_btn2 = st.columns([2, 1, 1])
                        with col_info:
                            st.info(f"Atualmente configurado para {st.session_state.edit_num_partes} partes envolvidas")
                        with col_btn1:
                            if st.button("➕ Adicionar Parte", key="edit_add_parte_btn"):
                                st.session_state.edit_num_partes += 1
                                st.rerun()
                        with col_btn2:
                            if st.button("➖ Remover Última", key="edit_remove_parte_btn") and st.session_state.edit_num_partes > 1:
                                st.session_state.edit_num_partes -= 1
                                st.rerun()
                        
                        st.markdown("---")

                        with st.form("edit_record_form"):
                            st.info(f"Editando registro de {record_type}")
                            updated_entries = {}
                            fields = FORM_DEFINITIONS.get(record_type, []) + COMMON_FIELDS

                            # Renderizar campos até "Partes Envolvidas"
                            partes_index = fields.index("Partes Envolvidas")
                            for field in fields[:partes_index]:
                                col_name = to_col_name(field)
                                current_value = record.get(col_name, "")
                                
                                if field == "Tipo de Ato":
                                    updated_entries[col_name] = st.selectbox(
                                        f"{field}:",
                                        options=TIPOS_DE_ATO,
                                        index=TIPOS_DE_ATO.index(current_value) if current_value in TIPOS_DE_ATO else None,
                                        key=f"edit_{col_name}"
                                    )
                                elif field in ["Data do Registro", "Data do Evento", "Data do Óbito"]:
                                    updated_entries[col_name] = st.text_input(
                                        f"{field}:", 
                                        value=current_value, 
                                        placeholder="DD/MM/AAAA",
                                        key=f"edit_{col_name}"
                                    )
                                else:
                                    updated_entries[col_name] = st.text_input(
                                        f"{field}:", 
                                        value=current_value, 
                                        key=f"edit_{col_name}"
                                    )

                            # Campos dinâmicos para "Partes Envolvidas"
                            st.markdown("#### Partes Envolvidas 👥")
                            edit_partes_inputs = []
                            for i in range(st.session_state.get('edit_num_partes', 1)):
                                val = partes_list[i] if i < len(partes_list) else ""
                                edit_partes_inputs.append(
                                    st.text_input(f"Parte Envolvida {i+1}", value=val, key=f"edit_parte_{i}")
                                )

                            # Campos restantes após "Partes Envolvidas"
                            st.markdown("---")
                            for field in fields[partes_index+1:]:
                                col_name = to_col_name(field)
                                current_value = record.get(col_name, "")
                                updated_entries[col_name] = st.text_input(
                                    f"{field}:", 
                                    value=current_value, 
                                    key=f"edit_{col_name}"
                                )

                            submitted = st.form_submit_button("Salvar Alterações")
                            
                            if submitted:
                                # Processar partes envolvidas
                                partes_values = [p.strip() for p in edit_partes_inputs if p.strip()]
                                updated_entries['partes_envolvidas'] = "; ".join(partes_values)

                                try:
                                    with engine.connect() as conn:
                                        now_utc = datetime.now(timezone.utc)
                                        
                                        set_clause = ", ".join([f"{col} = :{col}" for col in updated_entries.keys()])
                                        set_clause += ", ultima_alteracao_por = :user_email, atualizado_em = :now_utc"
                                        
                                        query = text(f"UPDATE registros SET {set_clause} WHERE id = :id")
                                        
                                        params = updated_entries
                                        params['id'] = record_id
                                        params['user_email'] = user_email
                                        params['now_utc'] = now_utc
                                        
                                        conn.execute(query, params)
                                        conn.commit()
                                        st.success("Registro atualizado com sucesso!")
                                        st.cache_data.clear()
                                        st.cache_resource.clear()
                                        del st.session_state.manage_action
                                        if 'edit_num_partes' in st.session_state: 
                                            del st.session_state.edit_num_partes
                                        st.rerun()
                                except Exception as e: 
                                    st.error(f"Ocorreu um erro ao atualizar: {e}")
                    else:
                        # Para outros tipos de registro (sem partes envolvidas dinâmicas)
                        with st.form("edit_record_form"):
                            st.info(f"Editando registro de {record_type}")
                            updated_entries = {}
                            fields = FORM_DEFINITIONS.get(record_type, []) + COMMON_FIELDS

                            for field in fields:
                                col_name = to_col_name(field)
                                current_value = record.get(col_name, "")
                                
                                if field == "Tipo de Ato":
                                    updated_entries[col_name] = st.selectbox(
                                        f"{field}:",
                                        options=TIPOS_DE_ATO,
                                        index=TIPOS_DE_ATO.index(current_value) if current_value in TIPOS_DE_ATO else None,
                                        key=f"edit_{col_name}"
                                    )
                                elif field in ["Data do Registro", "Data do Evento", "Data do Óbito"]:
                                    updated_entries[col_name] = st.text_input(
                                        f"{field}:", 
                                        value=current_value, 
                                        placeholder="DD/MM/AAAA",
                                        key=f"edit_{col_name}"
                                    )
                                else:
                                    updated_entries[col_name] = st.text_input(
                                        f"{field}:", 
                                        value=current_value, 
                                        key=f"edit_{col_name}"
                                    )

                            submitted = st.form_submit_button("Salvar Alterações")
                            
                            if submitted:
                                try:
                                    with engine.connect() as conn:
                                        now_utc = datetime.now(timezone.utc)
                                        
                                        set_clause = ", ".join([f"{col} = :{col}" for col in updated_entries.keys()])
                                        set_clause += ", ultima_alteracao_por = :user_email, atualizado_em = :now_utc"
                                        
                                        query = text(f"UPDATE registros SET {set_clause} WHERE id = :id")
                                        
                                        params = updated_entries
                                        params['id'] = record_id
                                        params['user_email'] = user_email
                                        params['now_utc'] = now_utc
                                        
                                        conn.execute(query, params)
                                        conn.commit()
                                        st.success("Registro atualizado com sucesso!")
                                        st.cache_data.clear()
                                        st.cache_resource.clear()
                                        del st.session_state.manage_action
                                        st.rerun()
                                except Exception as e: 
                                    st.error(f"Ocorreu um erro ao atualizar: {e}")

                elif action == "delete":
                    st.warning(f"Tem certeza que deseja excluir o registro ID {record_id}?")
                    if st.button("Confirmar Exclusão", type="primary"):
                        try:
                            with engine.connect() as conn:
                                conn.execute(text("DELETE FROM registros WHERE id = :id"), {'id': record_id})
                                conn.commit()
                                st.success("Registro excluído com sucesso!")
                                st.cache_data.clear()
                                st.cache_resource.clear()
                                del st.session_state.manage_action
                                del st.session_state.record_id
                                st.rerun()
                        except Exception as e: 
                            st.error(f"Erro ao excluir: {e}")
            
            # --- INÍCIO DO BLOCO CORRIGIDO E COM NOVA FUNCIONALIDADE ---
            st.markdown("---")
            st.subheader("Excluir Múltiplos Registros")
            st.warning("Esta funcionalidade permite excluir vários registros de uma vez. Use com cuidado!")

            # Inicializa o estado de controle no session_state
            if 'pending_multi_delete' not in st.session_state:
                st.session_state.pending_multi_delete = False
            if 'ids_to_delete_list' not in st.session_state:
                st.session_state.ids_to_delete_list = []

            ids_to_delete_input = st.text_input(
                "IDs para excluir (separados por vírgula ou em intervalo):",
                help="Use vírgulas para IDs avulsos (ex: 5, 8, 12) ou traço para intervalos (ex: 1-30). Pode combinar ambos (ex: 1-10, 15, 20-25).",
                key="multi_delete_input"
            )

            # Botão principal para INICIAR o processo de exclusão
            if st.button("Revisar para Exclusão Múltipla", key="review_delete_multiple_btn"):
                if not ids_to_delete_input:
                    st.error("Por favor, insira pelo menos um ID ou intervalo para excluir.")
                    st.session_state.pending_multi_delete = False # Garante que o estado seja falso
                else:
                    try:
                        final_ids = set()
                        parts = ids_to_delete_input.split(',')

                        for part in parts:
                            part = part.strip()
                            if not part:
                                continue

                            if '-' in part:
                                # Trata como um intervalo (range)
                                try:
                                    start_str, end_str = part.split('-')
                                    start = int(start_str.strip())
                                    end = int(end_str.strip())
                                    if start <= end:
                                        final_ids.update(range(start, end + 1))
                                except ValueError:
                                    # Ignora partes malformadas como '1-a' ou '1-2-3'
                                    st.warning(f"O intervalo '{part}' é inválido e será ignorado.")
                                    pass
                            elif part.isdigit():
                                # Trata como número único
                                final_ids.add(int(part))
                            else:
                                st.warning(f"A entrada '{part}' não é um ID válido e será ignorada.")
                        
                        id_list = sorted(list(final_ids))
                        
                        if not id_list:
                            st.error("Nenhum ID numérico válido foi encontrado na sua entrada.")
                            st.session_state.pending_multi_delete = False # Garante que o estado seja falso
                        else:
                            # ATIVA o modo de confirmação e salva a lista de IDs
                            st.session_state.pending_multi_delete = True
                            st.session_state.ids_to_delete_list = id_list
                            st.rerun() # Força um rerun para mostrar o painel de confirmação

                    except Exception as e:
                        st.error(f"Erro ao processar os IDs: {e}")
                        st.session_state.pending_multi_delete = False


            # Painel de confirmação que aparece SOMENTE se o estado 'pending_multi_delete' for True
            if st.session_state.get('pending_multi_delete', False):
                with st.expander("CONFIRMAR EXCLUSÃO MÚLTIPLA", expanded=True):
                    st.warning(f"Você está prestes a excluir {len(st.session_state.ids_to_delete_list)} registros. Esta ação é irreversível.")
                    
                    # O checkbox de confirmação
                    confirm = st.checkbox(f"Confirmo que desejo excluir PERMANENTEMENTE os registros com os IDs: {', '.join(map(str, st.session_state.ids_to_delete_list))}")
                    
                    col_confirm, col_cancel = st.columns(2)

                    with col_confirm:
                        # O botão de exclusão final agora funciona, pois seu estado depende do checkbox na mesma execução
                        if st.button("EXCLUIR AGORA", disabled=not confirm, type="primary"):
                            try:
                                with engine.connect() as conn:
                                    with conn.begin(): # Transação para segurança
                                        for record_id in st.session_state.ids_to_delete_list:
                                            conn.execute(text("DELETE FROM registros WHERE id = :id_val"), {"id_val": record_id})
                                
                                st.success(f"{len(st.session_state.ids_to_delete_list)} registros excluídos com sucesso!")
                                st.balloons()
                                
                                # Limpa os estados e recarrega a página
                                st.session_state.pending_multi_delete = False
                                st.session_state.ids_to_delete_list = []
                                st.cache_data.clear()
                                st.cache_resource.clear()
                                st.rerun()

                            except Exception as e:
                                st.error(f"Erro durante a exclusão final: {e}")

                    with col_cancel:
                        # Botão para cancelar a operação
                        if st.button("Cancelar"):
                            st.session_state.pending_multi_delete = False
                            st.session_state.ids_to_delete_list = []
                            st.rerun()
            # --- FIM DO BLOCO CORRIGIDO E COM NOVA FUNCIONALIDADE ---


    elif st.session_state.active_tab == "📤 Exportar Dados":
        st.header("Exportar Dados")
        if EXPORT_LIBS_AVAILABLE:
            all_books_export = get_distinct_values("fonte_livro")
            if not all_books_export: 
                st.warning("Nenhum registro encontrado para exportar.")
            else:
                selected_books_export = st.multiselect("Selecione os livros para exportar:", all_books_export, default=all_books_export)
                if selected_books_export:
                    export_format = st.radio("Formato de exportação:", ["Excel", "PDF"])
                    pdf_style = None
                    if export_format == "PDF":
                        st.subheader("Opções de PDF")
                        pdf_style = st.radio("Estilo do PDF:", ["Tabela (Índice/Catálogo)", "Relatório Detalhado"], help="**Tabela**: Visão geral compacta\n\n**Relatório Detalhado**: Todos os campos de cada registro")
                    if st.button("Gerar Arquivo para Download", type="primary"):
                        try:
                            with engine.connect() as conn:
                                query = text("SELECT * FROM registros WHERE fonte_livro = ANY(:books) ORDER BY tipo_registro, id")
                                result = conn.execute(query, {'books': selected_books_export})
                                all_records = [dict(row._mapping) for row in result]
                                if all_records:
                                    records_by_type = defaultdict(list)
                                    for record in all_records: 
                                        records_by_type[record['tipo_registro']].append(record)
                                    if export_format == "Excel":
                                        file_bytes = generate_excel_bytes(dict(records_by_type))
                                        if file_bytes: 
                                            st.download_button("📥 Baixar Excel", file_bytes, "cpindexator_export.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                                    else:
                                        if pdf_style == "Tabela (Índice/Catálogo)":
                                            file_bytes = generate_pdf_table(dict(records_by_type))
                                            filename = "cpindexator_indice.pdf"
                                        else:
                                            file_bytes = generate_pdf_detailed(dict(records_by_type))
                                            filename = "cpindexator_relatorio_detailed.pdf"
                                        if file_bytes: 
                                            st.download_button("📥 Baixar PDF", file_bytes, filename, "application/pdf")
                                else: 
                                    st.warning("Nenhum registro encontrado nos livros selecionados.")
                        except Exception as e: 
                            st.error(f"Erro ao gerar arquivo: {e}")
        else: 
            st.error("Bibliotecas de exportação não instaladas. Instale openpyxl e reportlab.")
    
    elif st.session_state.active_tab == "⚙️ Administração" and is_admin:
        st.header("⚙️ Administração do Banco de Dados")
        st.markdown("---")
        
        st.subheader("Alimentar Banco de Dados com Excel")
        st.info("Esta função permite adicionar múltiplos registros de um arquivo, definindo um Livro Fonte único para todos eles.")

        # Passo 1: Selecionar o tipo de registro
        record_type_upload = st.selectbox(
            "1. Selecione o Tipo de Registro para o upload:",
            list(FORM_DEFINITIONS.keys()),
            index=None,
            placeholder="Selecione o tipo...",
            key="upload_record_type"
        )

        # Passo 2: Fornecer o nome do livro fonte (NOVO CAMPO)
        book_source_upload = st.text_input(
            "2. Informe o Nome do Livro Fonte para este arquivo:",
            placeholder="Ex: Livro de Batismos 1880-1890",
            key="upload_book_source"
        )

        # Passo 3: Fazer o upload do arquivo
        uploaded_excel_file = st.file_uploader(
            "3. Escolha um arquivo Excel (.xlsx)",
            type="xlsx",
            key="excel_uploader"
        )

        if st.button("Iniciar Importação do Excel", type="primary"):
            # Validação Crítica: Verificar se todos os campos foram preenchidos
            book_name = book_source_upload.strip()
            if not record_type_upload or not book_name or not uploaded_excel_file:
                st.error("Erro: Todos os três campos (Tipo de Registro, Nome do Livro e Arquivo) são obrigatórios.")
                st.stop() # Interrompe a execução

            try:
                df = pd.read_excel(uploaded_excel_file, dtype=str).fillna('')
                
                # Renomeia as colunas do Excel para o formato do banco de dados
                original_columns = df.columns.tolist()
                column_mapping = {col: to_col_name(col) for col in original_columns}
                df.rename(columns=column_mapping, inplace=True)

                # Atribuição Forçada do Livro Fonte (A LÓGICA CHAVE)
                # Cria ou substitui a coluna 'fonte_livro' com o valor da interface.
                df['fonte_livro'] = book_name

                st.write("Pré-visualização dos dados a serem importados (Livro Fonte atribuído):", df.head())

                # Adiciona as colunas de metadados
                now_utc = datetime.now(timezone.utc)
                df['tipo_registro'] = record_type_upload
                df['criado_em'] = now_utc
                df['atualizado_em'] = now_utc
                df['criado_por'] = user_email
                df['ultima_alteracao_por'] = user_email
                
                # Remove colunas que não existem na tabela de destino para evitar erros
                db_cols = get_table_columns()
                df_filtered = df[[col for col in df.columns if col in db_cols]]

                # Salva os novos registros no banco de dados (modo 'append')
                with engine.connect() as conn:
                    with conn.begin(): # Usando uma transação
                        df_filtered.to_sql('registros', conn, if_exists='append', index=False)
                
                st.success(f"Importação concluída com sucesso! {len(df_filtered)} novos registros foram adicionados ao livro '{book_name}'.")
                st.balloons()
                st.cache_data.clear()
                st.cache_resource.clear()
                st.rerun()

            except Exception as e:
                st.error(f"Ocorreu um erro durante a importação do Excel: {e}")
                st.warning("Verifique se as colunas no arquivo Excel (exceto 'Fonte (Livro)') correspondem aos campos do formulário.")
        
        st.markdown("---")
        st.subheader("Gerenciar Livros")

        # Renomear Livro
        with st.expander("Renomear um Livro"):
            all_books_admin = get_distinct_values("fonte_livro")
            book_to_rename = st.selectbox("Livro de Origem", options=all_books_admin, index=None, key="rename_book_select")
            new_book_name = st.text_input("Novo Nome do Livro", key="new_book_name_input")
            
            if st.button("Renomear Livro", key="rename_book_btn"):
                if not book_to_rename or not new_book_name.strip():
                    st.warning("Selecione um livro de origem e digite um novo nome.")
                elif new_book_name.strip() in all_books_admin:
                    st.error(f"O nome '{new_book_name.strip()}' já existe. Escolha outro nome.")
                else:
                    try:
                        with engine.connect() as conn:
                            with conn.begin():
                                query = text("UPDATE registros SET fonte_livro = :new_name WHERE fonte_livro = :old_name")
                                conn.execute(query, {"new_name": new_book_name.strip(), "old_name": book_to_rename})
                        st.success(f"O livro '{book_to_rename}' foi renomeado para '{new_book_name.strip()}'.")
                        st.cache_data.clear()
                        st.cache_resource.clear()
                        st.rerun()
                    except Exception as e:
                        st.error(f"Erro ao renomear o livro: {e}")

        # Excluir Livro
        with st.expander("Excluir Registros de um Livro"):
            all_books_admin_del = get_distinct_values("fonte_livro")
            book_to_delete = st.selectbox("Livro a ser Excluído", options=all_books_admin_del, index=None, key="delete_book_select")
            
            if book_to_delete:
                confirm_delete_book = st.checkbox(f"Confirmo que desejo excluir PERMANENTEMENTE todos os registros do livro '{book_to_delete}'.", key="confirm_delete_book_check")
                if st.button("Excluir Livro Inteiro", disabled=not confirm_delete_book, type="primary"):
                    try:
                        with engine.connect() as conn:
                            with conn.begin():
                                query = text("DELETE FROM registros WHERE fonte_livro = :book_name")
                                result = conn.execute(query, {"book_name": book_to_delete})
                        st.success(f"Todos os {result.rowcount} registros do livro '{book_to_delete}' foram excluídos.")
                        st.cache_data.clear()
                        st.cache_resource.clear()
                        st.rerun()
                    except Exception as e:
                        st.error(f"Erro ao excluir os registros do livro: {e}")
        
        st.markdown("---")
        st.subheader("Backup e Restauração")
        
        # Exportar Backup CSV
        with st.expander("Exportar Backup Completo (CSV)"):
            st.info("Esta função exporta **todos** os registros da tabela para um arquivo CSV.")
            if st.button("Gerar Arquivo de Backup (CSV)"):
                try:
                    with engine.connect() as conn:
                        df = pd.read_sql_table('registros', conn)
                        csv = df.to_csv(index=False).encode('utf-8')
                        st.download_button("📥 Baixar Backup CSV", csv, "cpindexator_backup_completo.csv", "text/csv")
                except Exception as e: 
                    st.error(f"Erro ao exportar o banco de dados: {e}")

        # Importar de CSV (Substituir)
        with st.expander("Importar de um Backup (Substituir Tudo)"):
            st.warning("🚨 **Atenção:** A importação de CSV irá **APAGAR TODOS OS REGISTROS ATUAIS** antes de carregar os novos dados.")
            uploaded_file_csv = st.file_uploader("Escolha um arquivo CSV de backup", type="csv", key="csv_uploader")
            if uploaded_file_csv is not None:
                confirm_import_csv = st.checkbox("Confirmo que entendo que todos os dados atuais serão substituídos.")
                if st.button("Iniciar Importação do CSV", disabled=not confirm_import_csv):
                    try:
                        df_to_import = pd.read_csv(uploaded_file_csv)
                        with engine.connect() as conn:
                            with conn.begin():
                                conn.execute(text("DELETE FROM registros"))
                                df_to_import.to_sql('registros', conn, if_exists='append', index=False)
                        st.success(f"Importação concluída! {len(df_to_import)} registros importados.")
                        st.cache_data.clear()
                        st.cache_resource.clear()
                        st.rerun()
                    except Exception as e:
                        st.error(f"Erro durante a importação: {e}")
                        st.info("A operação foi revertida. Seus dados antigos estão seguros.")
    # --- FIM DA GRANDE MUDANÇA ---

# --- ROTEADOR PRINCIPAL ---
if 'user' not in st.session_state:
    st.session_state.user = None

if st.session_state.user is None:
    login_form()
else:
    main_app()
